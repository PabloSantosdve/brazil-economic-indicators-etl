"""Etapa Load do pipeline ETL de indicadores econômicos do Brasil.

Guarda a tabela final em arquivos que outras pessoas conseguem abrir.
Por enquanto, gera um relatório em Excel (.xlsx) formatado.
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Nome da aba do Excel.
NOME_ABA = "Indicadores Mensais"

# Formato de exibição de cada coluna. Os números continuam sendo números no
# Excel; o formato só muda como eles aparecem na tela.
# Atenção: IPCA e Selic já estão em "pontos percentuais" (1.25 significa 1,25%).
# Por isso o "%" vai entre aspas, como texto fixo. O formato de porcentagem do
# Excel multiplicaria por 100 e mostraria 125%.
FORMATOS = {
    "data": "mm/yyyy",
    "ipca": '0.00"%"',
    "selic": '0.00"%"',
    "dolar": '"R$" #,##0.0000',
}

# Largura mínima de cada coluna, em caracteres.
LARGURA_MINIMA = 14


def salvar_excel(df: pd.DataFrame, caminho: str) -> None:
    """Salva a tabela mensal final em um arquivo Excel formatado.

    Args:
        df: tabela final do Transform, com o mês no índice (chamado 'data') e
            as colunas 'ipca', 'selic' e 'dolar'.
        caminho: onde salvar o arquivo, por exemplo "data/indicadores_mensais.xlsx".
            A pasta é criada automaticamente se não existir.
    """
    # Cria a pasta de destino se ela ainda não existir (ex.: "data/").
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)

    # A data está no índice da tabela. O reset_index a transforma em uma coluna
    # comum, que é o que vai para o Excel. O rename_axis garante o nome 'data'.
    tabela = df.rename_axis("data").reset_index()

    # Estilos que serão reaproveitados em várias células.
    fonte_cabecalho = Font(bold=True, color="FFFFFF")
    fundo_cabecalho = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    lado = Side(style="thin", color="BFBFBF")
    borda_fina = Border(left=lado, right=lado, top=lado, bottom=lado)

    # O ExcelWriter abre o arquivo, e o bloco "with" salva e fecha tudo no fim.
    # Toda a formatação precisa acontecer dentro dele.
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        # Escreve os dados. O index=False evita uma coluna extra de numeração.
        tabela.to_excel(writer, sheet_name=NOME_ABA, index=False)
        aba = writer.sheets[NOME_ABA]

        # Cabeçalho (linha 1): fonte branca em negrito, fundo azul, centralizado.
        for celula in aba[1]:
            celula.font = fonte_cabecalho
            celula.fill = fundo_cabecalho
            celula.alignment = Alignment(horizontal="center", vertical="center")
            celula.border = borda_fina

        # Dados (da linha 2 em diante): borda, alinhamento e formato numérico.
        for linha in aba.iter_rows(min_row=2, max_row=aba.max_row):
            for celula in linha:
                nome_coluna = tabela.columns[celula.column - 1]
                celula.border = borda_fina
                celula.number_format = FORMATOS.get(nome_coluna, "General")
                # Data à esquerda, números à direita.
                alinhamento = "left" if celula.column == 1 else "right"
                celula.alignment = Alignment(horizontal=alinhamento)

        # Largura das colunas. Usamos o tamanho do título, com um mínimo fixo,
        # porque o tamanho do conteúdo não reflete o que aparece na tela
        # (o Excel mostra 4 casas decimais, mas o número guardado tem mais).
        for posicao, nome in enumerate(tabela.columns, start=1):
            letra = get_column_letter(posicao)
            aba.column_dimensions[letra].width = max(len(str(nome)) + 4, LARGURA_MINIMA)